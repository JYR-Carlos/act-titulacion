"""
calcular_tsr.py — MÉTRICA 4 del MVP: Task Success Rate (Sección 11).

Lee la planilla de las pruebas con usuarios y calcula el % de éxito, su
intervalo de confianza y el veredicto contra el umbral del 80%, **destacado en
rojo si no se alcanza**. Verifica además que el protocolo se haya respetado
(mínimo 10 pruebas con participantes distintos).

Uso:
    python scripts/calcular_tsr.py                                  # data/tsr/planilla_tsr.csv
    python scripts/calcular_tsr.py --planilla data/tsr/planilla_tsr.csv
    python scripts/calcular_tsr.py --planilla planilla.xlsx         # requiere openpyxl

La planilla en blanco está en `data/tsr/planilla_tsr_plantilla.csv` y el guion
del escenario en `docs/PROTOCOLO_TSR.md`.

Columnas esperadas (las demás se ignoran):
    participante          identificador anonimizado (P01, P02, ...) — obligatorio
    secuencia_id          qué guion de 3 señas se ejecutó (S1..S5)
    senas_ejecutadas      las 3 glosas, separadas por " + "
    interpretacion        lo que el participante dijo que entendió
    exito                 si / no  (también acepta 1/0, true/false, x)
    comentarios           feedback cualitativo (legibilidad, utilidad)
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

import _raiz  # noqa: F401  (raíz del repo en sys.path; debe ir antes que lsch_mr)
from lsch_mr import config
from lsch_mr import metricas_tsr as tsr
from lsch_mr.consola import configurar_utf8, amarillo, negrita, rojo, verde

configurar_utf8()

# Cómo se interpreta la columna `exito`. Se acepta un abanico amplio porque la
# planilla la rellena una persona a mano, a veces en Excel y en español.
_VERDADEROS = {"si", "sí", "s", "yes", "y", "1", "true", "verdadero", "x", "ok",
               "exito", "éxito"}
_FALSOS = {"no", "n", "0", "false", "falso", "fallo", "fracaso", ""}


def _a_bool(valor, fila: int) -> bool:
    """Convierte la columna `exito` a bool. Un valor no reconocido es un error.

    Deliberadamente estricto: interpretar en silencio un "tal vez" o un "parcial"
    como fallo (o como éxito) falsearía la métrica sin dejar rastro. Mejor parar
    y que la persona decida.
    """
    if isinstance(valor, bool):
        return valor
    texto = str(valor if valor is not None else "").strip().lower()
    if texto in _VERDADEROS:
        return True
    if texto in _FALSOS:
        return False
    raise ValueError(
        f"Fila {fila}: no entiendo el valor de 'exito' = {valor!r}. "
        f"Usa 'si' o 'no' (el éxito es binario por definición de la métrica: "
        "interpretación correcta SIN comunicación adicional).")


def _leer_csv(ruta: Path) -> list[dict]:
    with ruta.open(encoding="utf-8-sig") as f:   # -sig: Excel escribe BOM
        return list(csv.DictReader(f))


def _leer_xlsx(ruta: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit(rojo(
            f"Para leer {ruta.name} hace falta openpyxl:  pip install openpyxl\n"
            "  (o guarda la planilla como CSV desde Excel y pásala con --planilla)"))
    hoja = load_workbook(ruta, data_only=True).active
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return []
    encabezados = [str(c).strip() if c is not None else "" for c in filas[0]]
    return [dict(zip(encabezados, fila)) for fila in filas[1:]
            if any(c is not None and str(c).strip() for c in fila)]


def cargar_registros(ruta: Path) -> list[dict]:
    """Lee la planilla (CSV o XLSX) y normaliza las columnas que importan."""
    crudas = (_leer_xlsx(ruta) if ruta.suffix.lower() in (".xlsx", ".xlsm")
              else _leer_csv(ruta))

    registros = []
    for i, r in enumerate(crudas, start=2):     # +2: fila 1 = encabezados
        normalizada = {str(k).strip().lower(): v for k, v in r.items() if k}
        participante = str(normalizada.get("participante", "") or "").strip()
        # Fila totalmente vacía (habitual al final de una planilla de Excel).
        if not participante and not str(normalizada.get("exito", "") or "").strip():
            continue
        registros.append({
            "fila": i,
            "participante": participante,
            "secuencia_id": str(normalizada.get("secuencia_id", "") or "").strip(),
            "senas_ejecutadas": str(normalizada.get("senas_ejecutadas", "") or "").strip(),
            "interpretacion": str(normalizada.get("interpretacion", "") or "").strip(),
            "exito": _a_bool(normalizada.get("exito"), i),
            "comentarios": str(normalizada.get("comentarios", "") or "").strip(),
        })
    return registros


def _imprimir(resumen: dict, registros: list[dict]) -> None:
    print()
    print(negrita("=" * 74))
    print(negrita("  MÉTRICA 4 — TASK SUCCESS RATE (pruebas con usuarios)"))
    print(negrita("=" * 74))
    print(f"  pruebas registradas   : {resumen['n_pruebas']}")
    print(f"  participantes distintos: {resumen['n_participantes']}")
    print(f"  éxitos / fallos       : {resumen['n_exitos']} / {resumen['n_fallos']}")

    print()
    linea = (f"  TSR: {resumen['tsr']:.1%}   "
             f"(objetivo MVP >= {resumen['objetivo']:.0%})")
    if resumen["cumple"]:
        print(verde(linea + "   -> CUMPLE"))
    elif resumen["cumple_umbral_sin_validar_protocolo"]:
        print(amarillo(linea + "   -> alcanza el % pero NO cumple el protocolo"))
    else:
        print(rojo(linea + "   -> NO CUMPLE"))

    ic = resumen["ic95"]
    print(f"  IC 95% (Wilson)       : [{ic[0]:.1%}, {ic[1]:.1%}]")
    if resumen["n_pruebas"] and not resumen["ic95_respalda_objetivo"]:
        print(amarillo(
            "  [!] El límite inferior del intervalo queda por debajo del objetivo: "
            f"con n={resumen['n_pruebas']}\n      la muestra es compatible con un "
            "sistema que no alcanza el 80% real. Reporta el\n      intervalo junto "
            "al porcentaje, no el porcentaje solo."))

    for alerta in resumen["alertas"]:
        print(rojo(f"  [PROTOCOLO] {alerta}"))

    # -- Desglose por secuencia -------------------------------------------- #
    desglose = tsr.desglose_por_secuencia(registros)
    if len(desglose) > 1:
        print()
        print(negrita("  Por secuencia"))
        for d in desglose:
            marca = "" if d["tsr"] >= resumen["objetivo"] else "  <-- concentra fallos"
            print(f"    {d['secuencia_id']:<10} {d['exitos']}/{d['n']}"
                  f"   {d['tsr']:.0%}{marca}")

    # -- Fallos, con lo que el participante entendió ------------------------ #
    fallos = [r for r in registros if not r["exito"]]
    if fallos:
        print()
        print(negrita("  Pruebas fallidas"))
        for r in fallos:
            print(f"    {r['participante']:<6} [{r['secuencia_id']}] "
                  f"ejecutado: {r['senas_ejecutadas'] or '(sin registrar)'}")
            print(f"           interpretó: {r['interpretacion'] or '(sin registrar)'}")

    # -- Feedback cualitativo ----------------------------------------------- #
    comentarios = tsr.comentarios_cualitativos(registros)
    if comentarios:
        print()
        print(negrita("  Feedback cualitativo (legibilidad / utilidad)"))
        for c in comentarios:
            marca = "OK  " if c["exito"] else "FALLO"
            print(f"    [{marca}] {c['participante']}: {c['comentarios']}")
    else:
        print()
        print(amarillo("  [!] Ningún comentario cualitativo registrado. El diseño "
                       "pide recoger\n      percepción de legibilidad y utilidad "
                       "además del porcentaje."))
    print()


def main() -> int:
    ap = argparse.ArgumentParser(
        description="Task Success Rate de las pruebas con usuarios (métrica 4)")
    ap.add_argument("--planilla",
                    default=str(config.DATA_TSR_DIR / "planilla_tsr.csv"),
                    help="CSV o XLSX con los registros de las pruebas")
    ap.add_argument("--objetivo", type=float, default=tsr.TSR_OBJETIVO)
    ap.add_argument("--min-pruebas", type=int, default=tsr.MIN_PRUEBAS)
    ap.add_argument("--salida",
                    default=str(config.OUTPUTS_REPORTS_DIR / "tsr.json"))
    args = ap.parse_args()

    ruta = Path(args.planilla)
    if not ruta.exists():
        print(rojo(f"[tsr] No existe {ruta}."))
        print("  Copia la plantilla en blanco y rellénala durante las pruebas:")
        print(r"    copy data\tsr\planilla_tsr_plantilla.csv data\tsr\planilla_tsr.csv")
        print("  El guion del escenario está en docs/PROTOCOLO_TSR.md")
        return 1

    try:
        registros = cargar_registros(ruta)
    except ValueError as exc:
        print(rojo(f"[tsr] {exc}"))
        return 1

    if not registros:
        print(rojo(f"[tsr] {ruta} no tiene ninguna prueba registrada."))
        return 1

    resumen = tsr.calcular_tsr(registros, objetivo=args.objetivo,
                               min_pruebas=args.min_pruebas)
    _imprimir(resumen, registros)

    salida = Path(args.salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    salida.write_text(json.dumps({
        "metrica": "4 - task success rate",
        "planilla": str(ruta),
        "resumen": resumen,
        "por_secuencia": tsr.desglose_por_secuencia(registros),
        "comentarios": tsr.comentarios_cualitativos(registros),
        "registros": registros,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"  reporte: {salida}")
    print()
    return 0 if resumen["cumple"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
